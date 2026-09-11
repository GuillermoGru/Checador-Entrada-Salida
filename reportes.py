# reportes.py

# ------Importaciones------
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from datetime import datetime

# ------Herramientas Matemáticas y de Tiempo------

def sumar_horas(lista_horas):
    # Agarra una lista de textos como "08:30" y "01:15", los parte a la mitad, 
    # suma todo en minutos para que no falle y luego lo vuelve a armar como reloj.
    total_minutos = 0
    for t in lista_horas:
        if pd.isna(t) or t in ['0', 0, '-1', -1, '']:
            continue
        try:
            h, m = map(int, str(t).split(':'))
            total_minutos += h * 60 + m
        except:
            pass
            
    if total_minutos == 0: return ""
    return f"{total_minutos // 60:02d}:{total_minutos % 60:02d}"

def calc_dif_mins(actual, official, is_in=True):
    # Compara a qué hora llegó contra a qué hora debía llegar.
    actual_str = str(actual).replace('nan', '').strip()[:5]
    official_str = str(official).replace('nan', '').strip()[:5]
    
    if not actual_str or not official_str: return 0
        
    try:
        t_act = datetime.strptime(actual_str, '%H:%M')
        t_off = datetime.strptime(official_str, '%H:%M')
        diff = (t_act - t_off).total_seconds()
        
        # Arreglamos la matemática si brincamos la medianoche
        if diff < -12 * 3600: diff += 24 * 3600
        elif diff > 12 * 3600: diff -= 24 * 3600
        
        # Si es la salida, la lógica es al revés (salir antes de tiempo es lo malo)
        if not is_in: diff = -diff
            
        if diff > 0: return int(diff // 60) # Regresamos todo convertido a minutos
        return 0
    except:
        return 0

# ------Gestor de Observaciones y Permisos ------

def buscar_observacion_dia(df_permisos, emp_id, fecha_dt):
    # Busca si hay un permiso en la fecha especificada
    if df_permisos is None or df_permisos.empty: return ""
    try:
        emp_id_str = str(emp_id).replace('.0', '').strip()
        if 'ID del Empleado' not in df_permisos.columns: return ""
            
        df_fil = df_permisos[df_permisos['ID del Empleado'].astype(str).str.replace(r'\.0$', '', regex=True).str.strip() == emp_id_str]
        if df_fil.empty: return ""
        
        for _, row in df_fil.iterrows():
            
            # Intentamos sacar la fecha inicial y final del Excel
            p_ini = pd.to_datetime(row.get('Fecha y Hora Inicial'), errors='coerce')
            p_fin = pd.to_datetime(row.get('Fecha y Hora Final'), errors='coerce')
            
            # Si al menos tiene fecha inicial, podemos usarlo
            if pd.isna(p_ini): continue
            
            # Si no pusieron fecha final en el Excel, asumimos que es permiso de un solo día
            if pd.isna(p_fin): p_fin = p_ini 
            
            # Normalizamos (quitamos las horas) para comparar solo los días completos
            d_ini = p_ini.normalize()
            d_fin = p_fin.normalize()
            d_eval = fecha_dt.normalize()
            
            # Verificamos si el día que estamos evaluando entra en el rango del permiso
            if d_ini <= d_eval <= d_fin:
                # Extraemos estrictamente la Categoría de Permiso
                categoria = str(row.get('Categoría de Permiso', ''))
                if categoria.lower() == 'nan' or not categoria.strip(): 
                    categoria = "Permiso Registrado"
                return categoria
    except:
        pass
    return ""

def buscar_incidencia_dia(df_incidencias, fecha_dt):
    # Busca si hay una incidencia en el calendario para la fecha especificada
    if df_incidencias is None or df_incidencias.empty or 'Fecha_dt' not in df_incidencias.columns: 
        return ""
    
    d_eval = fecha_dt.normalize()
    fila = df_incidencias[df_incidencias['Fecha_dt'] == d_eval]
    
    if not fila.empty:
        tipo = str(fila.iloc[0].get('Tipo', '')).strip()
        if tipo.lower() != 'nan' and tipo:
            return tipo
    return ""

# ------Gestor de Fechas y Periodos------

def segmentar_fechas(rango_fechas, tipo_corte):
    # Agarra el mes entero y lo divide según lo que pidió el usuario (semanas, quincenas...)
    chunks = []
    chunk_actual = []
    
    for fecha in rango_fechas:
        if not chunk_actual:
            chunk_actual.append(fecha)
            continue
            
        fecha_anterior = chunk_actual[-1]
        cortar = False
        
        # Aquí le decimos cuándo cortar la tabla
        if tipo_corte == 'Semanal':
            if fecha.weekday() == 0 and fecha_anterior.weekday() == 6: cortar = True
        elif tipo_corte == 'Quincenal':
            if (fecha_anterior.day <= 15 and fecha.day > 15) or (fecha_anterior.month != fecha.month): cortar = True
        elif tipo_corte == 'Mensual':
            if fecha_anterior.month != fecha.month: cortar = True
                
        if cortar:
            chunks.append(chunk_actual)
            chunk_actual = [fecha]
        else:
            chunk_actual.append(fecha)
            
    if chunk_actual: chunks.append(chunk_actual)
    return chunks

# ------Motor Principal de Reportes------

def _generar_reporte_base(ruta_bd: Path, fecha_inicio: str, fecha_fin: str, tipo_corte: str, es_velador: bool):
    #Lee si es velador o empleado normal y adapta todo el proceso.
    hoja_datos = 'DatosVeladores' if es_velador else 'Datos'
    tipo_nombre = 'Veladores' if es_velador else 'Empleados'
    
    # Abrimos todas las páginas que necesitamos de A2026.xlsx
    try:
        dfs = pd.read_excel(ruta_bd, sheet_name=None, header=1)
        if 'Empleados' not in dfs or hoja_datos not in dfs:
            return None
            
        df_emp = dfs['Empleados'].fillna('')
        df_datos = dfs[hoja_datos].dropna(subset=['ID de Empleado', 'Fecha']).copy()
        df_hor = dfs.get('Horarios')
            
        try: df_permisos = pd.read_excel(ruta_bd, sheet_name='Permisos', header=0)
        except: df_permisos = pd.DataFrame()
        
        try: 
            df_incidencias = pd.read_excel(ruta_bd, sheet_name='Incidencias', header=1)
            if not df_incidencias.empty and 'Fecha' in df_incidencias.columns:
                df_incidencias['Fecha_dt'] = pd.to_datetime(df_incidencias['Fecha'], errors='coerce').dt.normalize()
        except: 
            df_incidencias = pd.DataFrame()
            
    except Exception:
        return None

    # Vemos qué tan estrictos somos con los retardos leyendo la variable de Excel
    rango_horarios = 30
    try:
        df_vars = pd.read_excel(ruta_bd, sheet_name='Variables', header=None)
        fila_rango = df_vars[df_vars[0].astype(str).str.contains('Rango de horarios', case=False, na=False)]
        if not fila_rango.empty:
            val = pd.to_numeric(fila_rango[1].iloc[0], errors='coerce')
            if pd.notna(val): rango_horarios = int(val)
    except:
        pass

    if df_hor is not None:
        df_hor['ID de Empleado'] = df_hor['ID de Empleado'].astype(str).str.replace(r'\.0$', '', regex=True)

    # Cortamos los datos para solo ver el rango de fechas que pidió el usuario
    df_datos['Fecha_dt'] = pd.to_datetime(df_datos['Fecha'], errors='coerce')
    try:
        dt_inicio = pd.to_datetime(fecha_inicio, format="%d/%m/%Y")
        dt_fin = pd.to_datetime(fecha_fin, format="%d/%m/%Y")
    except:
        return None

    mask = (df_datos['Fecha_dt'] >= dt_inicio) & (df_datos['Fecha_dt'] <= dt_fin)
    df_rango = df_datos[mask].copy()
    
    if df_rango.empty:
        return None

    # Preparamos para dibujar en Excel
    rango_fechas = pd.date_range(start=dt_inicio, end=dt_fin)
    chunks_fechas = segmentar_fechas(rango_fechas, tipo_corte)
    dias_es = {0: 'Lun', 1: 'Mar', 2: 'Mie', 3: 'Jue', 4: 'Vie', 5: 'Sab', 6: 'Dom'}
    ids_en_periodo = df_rango['ID de Empleado'].unique()
    
    wb = Workbook()
    ws = wb.active
    ws.title = f"Reporte {tipo_nombre}"

    fuente_titulo = Font(name='Arial', size=11, bold=True)
    fuente_normal = Font(name='Arial', size=11)
    alineacion_centro = Alignment(horizontal='center', vertical='center', wrap_text=True)
    borde_delgado = Border(
        left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin')
    )

    fila_actual = 1
    ws.cell(row=fila_actual, column=1, value=f"Periodo: {fecha_inicio} - {fecha_fin}  |  Corte: {tipo_corte}").font = fuente_titulo
    fila_actual += 2
    cabeceras = ['Fecha', 'Día', 'Horario Entrada', 'Horario Salida', 'Horas', 'Dif. Ent', 'Dif. Sal', 'Incumplió', 'Observación', 'Total']
    
    # Dibujamos empleado por empleado
    for emp_id in sorted(ids_en_periodo):
        emp_id_str = str(emp_id).replace('.0', '')
        
        emp_info = df_emp[df_emp['ID de Empleado'].astype(str).str.replace(r'\.0$', '', regex=True) == emp_id_str]
        nombre_completo = f"{emp_id_str}  {emp_info.iloc[0]['Apellido']} {emp_info.iloc[0]['Nombre']}" if not emp_info.empty else f"{emp_id_str}  (Desconocido)"
        ws.cell(row=fila_actual, column=1, value=nombre_completo).font = fuente_titulo
        fila_actual += 1
        
        # Rescatamos el horario oficial de este trabajador
        h_in1, h_out1, h_in2, h_out2 = "", "", "", ""
        if df_hor is not None:
            emp_hor = df_hor[df_hor['ID de Empleado'] == emp_id_str]
            if not emp_hor.empty:
                h_in1 = str(emp_hor.iloc[0].get('Hora inicio', '')).replace('nan', '')
                h_out1 = str(emp_hor.iloc[0].get('Hora final', '')).replace('nan', '')
                h_in2 = str(emp_hor.iloc[0].get('Hora inicio2', '')).replace('nan', '')
                h_out2 = str(emp_hor.iloc[0].get('Hora final2', '')).replace('nan', '')
        
        df_emp_datos = df_rango[df_rango['ID de Empleado'] == emp_id]
        
        # Dibujamos las semanas o quincenas
        for bloque_fechas in chunks_fechas:
            for col_idx, cabecera in enumerate(cabeceras, 1):
                celda = ws.cell(row=fila_actual, column=col_idx, value=cabecera)
                celda.font = fuente_titulo
                celda.alignment = alineacion_centro
                celda.border = borde_delgado
            fila_actual += 1
            
            total_horas_periodo, total_dif_ent_periodo, total_dif_sal_periodo = [], [], []
            total_incumplio_periodo = 0
            
            # Revisamos día por día en el calendario
            for fecha_eval in bloque_fechas:
                marcas_dia = df_emp_datos[df_emp_datos['Fecha_dt'] == fecha_eval]
                fecha_str = fecha_eval.strftime("%d/%m/%Y")
                dia_str = dias_es[fecha_eval.weekday()]
                
                # Buscamos permisos e incidencias ANTES de ver si vino o no
                texto_permiso = buscar_observacion_dia(df_permisos, emp_id_str, fecha_eval)
                texto_incidencia = buscar_incidencia_dia(df_incidencias, fecha_eval)
                
                # Si no vino a trabajar, dejamos la línea en blanco pero evaluamos la observación
                if marcas_dia.empty:
                    observaciones_dia_vacio = []
                    
                    # El orden de los .append dicta cómo se imprime el texto
                    if texto_incidencia:
                        observaciones_dia_vacio.append(texto_incidencia)
                    if texto_permiso:
                        observaciones_dia_vacio.append(texto_permiso)
                    
                    # Unimos usando un espacio para que quede: "Efeméride Permiso Autorizado"
                    texto_obs_final = " ".join(observaciones_dia_vacio)

                    for col_idx in range(1, 11):
                        celda = ws.cell(row=fila_actual, column=col_idx)
                        celda.border = borde_delgado
                        celda.font = fuente_normal
                        celda.alignment = alineacion_centro
                        if col_idx == 1: celda.value = fecha_str
                        if col_idx == 2: celda.value = dia_str
                        if col_idx == 9: celda.value = texto_obs_final 
                    fila_actual += 1
                else:
                    # Si sí vino, revisamos cuántas veces pasó la tarjeta
                    total_horas_dia = []
                    lista_marcas = marcas_dia.to_dict('records')
                    num_marcas = len(lista_marcas)
                    valid_index = 0
                    
                    for idx, marca in enumerate(lista_marcas):
                        dif_ent_val, dif_sal_val, incumplio_val = "", "", ""
                        
                        observaciones_renglon = []
                        # Solo ponemos la observación completa en la primera checada del día
                        if idx == 0:
                            if texto_incidencia:
                                observaciones_renglon.append(texto_incidencia)
                            if texto_permiso:
                                observaciones_renglon.append(texto_permiso)
                        
                        # Unimos usando un salto de línea
                        texto_obs_final = "\n".join(observaciones_renglon)
                        
                        horas_str = str(marca.get('Diferencia horas', '')).replace('nan', '')
                        act_in = str(marca.get('Hora entrada', '')).replace('nan', '')
                        act_out = str(marca.get('Hora salida', '')).replace('nan', '')
                        
                        # Analizamos retardos solo en turnos reales (ignoramos dobles checadas por accidente)
                        if horas_str not in ['0', '-1', 0, -1, '']:
                            if valid_index == 0: off_in, off_out = h_in1, h_out1
                            elif valid_index == 1: off_in, off_out = h_in2, h_out2
                            else: off_in, off_out = "", ""
                                
                            valid_index += 1
                            mins_ent = calc_dif_mins(act_in, off_in, is_in=True)
                            mins_sal = calc_dif_mins(act_out, off_out, is_in=False)
                            
                            if off_in:
                                dif_ent_val = f"{mins_ent//60:02d}:{mins_ent%60:02d}"
                                if mins_ent > 0: total_dif_ent_periodo.append(dif_ent_val)
                                
                            if off_out:
                                dif_sal_val = f"{mins_sal//60:02d}:{mins_sal%60:02d}"
                                if mins_sal > 0: total_dif_sal_periodo.append(dif_sal_val)
                                
                            # Castigamos si se pasa del margen de tolerancia
                            if mins_ent > rango_horarios or mins_sal > rango_horarios:
                                incumplio_val = 1
                                # El perdón se otorga si tuvo permiso en ese día
                                if texto_permiso:
                                    incumplio_val = ""
                                else:
                                    total_incumplio_periodo += 1
                                    
                        # Plasmamos toda esta matemática en las celdas del Excel
                        for col_idx in range(1, 11):
                            celda = ws.cell(row=fila_actual, column=col_idx)
                            celda.border = borde_delgado
                            celda.font = fuente_normal
                            celda.alignment = alineacion_centro
                            
                            if idx == 0:
                                if col_idx == 1: celda.value = fecha_str
                                if col_idx == 2: celda.value = dia_str
                                
                            if col_idx == 3: celda.value = act_in
                            if col_idx == 4: celda.value = act_out
                            if col_idx == 5: 
                                celda.value = horas_str
                                total_horas_dia.append(horas_str)
                                total_horas_periodo.append(horas_str)
                            
                            if col_idx == 6: celda.value = dif_ent_val
                            if col_idx == 7: celda.value = dif_sal_val
                            if col_idx == 8: celda.value = incumplio_val
                            if col_idx == 9: celda.value = texto_obs_final # Insertamos la variable concatenada
                            
                            if col_idx == 10 and idx == num_marcas - 1:
                                celda.value = sumar_horas(total_horas_dia)
                                celda.font = fuente_titulo 
                                
                        fila_actual += 1
                        
            # Imprimimos los totales al final de la semana/quincena
            for col_idx in range(1, 11):
                celda = ws.cell(row=fila_actual, column=col_idx)
                celda.border = borde_delgado
                if col_idx == 6:
                    celda.value = sumar_horas(total_dif_ent_periodo)
                    celda.font = fuente_titulo
                    celda.alignment = alineacion_centro
                if col_idx == 7:
                    celda.value = sumar_horas(total_dif_sal_periodo)
                    celda.font = fuente_titulo
                    celda.alignment = alineacion_centro
                if col_idx == 8:
                    celda.value = total_incumplio_periodo if total_incumplio_periodo > 0 else ""
                    celda.font = fuente_titulo
                    celda.alignment = alineacion_centro
                if col_idx == 9:
                    celda.value = "Total"
                    celda.font = fuente_titulo
                    celda.alignment = alineacion_centro
                if col_idx == 10:
                    celda.value = sumar_horas(total_horas_periodo)
                    celda.font = fuente_titulo
                    celda.alignment = alineacion_centro
            
            fila_actual += 2 
        fila_actual += 2 

    # Hacemos más anchas las columnas para que quepa bien el texto
    anchos = [12, 6, 15, 15, 8, 10, 10, 12, 28, 10]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[chr(64+i)].width = ancho

    # Guardamos el archivo y avisamos que terminamos
    nombre_archivo = f"Reporte_{tipo_nombre}_{tipo_corte}_{fecha_inicio.replace('/','-')}_al_{fecha_fin.replace('/','-')}.xlsx"
    wb.save(nombre_archivo)
    return nombre_archivo

# ------Interruptores de Reporte------

def generar_reporte_empleados(ruta_bd: Path, fecha_inicio: str, fecha_fin: str, tipo_corte: str):
    return _generar_reporte_base(ruta_bd, fecha_inicio, fecha_fin, tipo_corte, es_velador=False)

def generar_reporte_veladores(ruta_bd: Path, fecha_inicio: str, fecha_fin: str, tipo_corte: str):
    return _generar_reporte_base(ruta_bd, fecha_inicio, fecha_fin, tipo_corte, es_velador=True)